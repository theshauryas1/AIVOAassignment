"""
generate_sample_docs.py – Generate realistic pharmaceutical API deviation documents.
Creates:
  1. sample_documents/deviation_metformin_oos.pdf
  2. sample_documents/deviation_email_incident.txt
  3. sample_documents/deviation_batch_excursion.docx
"""
import os
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import docx

os.makedirs("sample_documents", exist_ok=True)

# ── 1. Create PDF ───────────────────────────────────────────────────────────
pdf_path = "sample_documents/deviation_metformin_oos.pdf"
doc = SimpleDocTemplate(
    pdf_path,
    pagesize=letter,
    rightMargin=36,
    leftMargin=36,
    topMargin=36,
    bottomMargin=36
)

styles = getSampleStyleSheet()
title_style = ParagraphStyle(
    'DocTitle',
    parent=styles['Heading1'],
    fontSize=16,
    textColor=colors.HexColor('#0f172a'),
    spaceAfter=6,
)
subtitle_style = ParagraphStyle(
    'DocSubTitle',
    parent=styles['Normal'],
    fontSize=10,
    textColor=colors.HexColor('#64748b'),
    spaceAfter=14,
)
body_style = ParagraphStyle(
    'DocBody',
    parent=styles['Normal'],
    fontSize=9.5,
    leading=14,
    textColor=colors.HexColor('#1e293b'),
)
bold_label = ParagraphStyle(
    'BoldLabel',
    parent=styles['Normal'],
    fontSize=9.5,
    fontName='Helvetica-Bold',
    textColor=colors.HexColor('#0f172a'),
)

story = []

# Header
story.append(Paragraph("VASUDHA PHARMA CHEM LIMITED – QUALITY ASSURANCE DIVISION", subtitle_style))
story.append(Paragraph("INTERNAL DEVIATION & NON-CONFORMANCE NOTIFICATION", title_style))
story.append(Spacer(1, 10))

# Table with structured metadata
data = [
    [Paragraph("<b>Deviation Number:</b>", bold_label), Paragraph("DEV-2024-0814-A01", body_style),
     Paragraph("<b>Date of Occurrence:</b>", bold_label), Paragraph("14-08-2024", body_style)],
    [Paragraph("<b>Site / Plant:</b>", bold_label), Paragraph("API Manufacturing Unit - Block B", body_style),
     Paragraph("<b>Source of Event:</b>", bold_label), Paragraph("Manufacturing (In-Process Control)", body_style)],
    [Paragraph("<b>Product / Material:</b>", bold_label), Paragraph("Metformin Hydrochloride API (BP/USP)", body_style),
     Paragraph("<b>Batch / Lot Number:</b>", bold_label), Paragraph("MFH260712A", body_style)],
    [Paragraph("<b>Equipment Involved:</b>", bold_label), Paragraph("Glass-Lined Reactor R-102", body_style),
     Paragraph("<b>Initial Severity:</b>", bold_label), Paragraph("High (Process Parameter OOS)", body_style)],
]

t = Table(data, colWidths=[120, 150, 120, 150])
t.setStyle(TableStyle([
    ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
    ('TOPPADDING', (0,0), (-1,-1), 6),
    ('BOTTOMPADDING', (0,0), (-1,-1), 6),
]))
story.append(t)
story.append(Spacer(1, 14))

# Event Description
story.append(Paragraph("<b>1. Detailed Incident Description:</b>", bold_label))
story.append(Paragraph(
    "During Stage-II crystallization of Metformin Hydrochloride synthesis in Glass-Lined Reactor R-102 "
    "(Batch No: MFH260712A), the automated chilled brine supply valve (CV-102B) malfunctioned and became locked in the 15% open position. "
    "This caused the internal crystallization temperature to rise from the validated target range of 20.0°C – 24.0°C to an excursion peak of 48.5°C. "
    "The out-of-specification temperature condition persisted for approximately 42 minutes before the plant operator noticed the local alarm and initiated manual bypass override. "
    "The crystallization process was subsequently completed under manual cooling control.",
    body_style
))
story.append(Spacer(1, 10))

# Immediate actions
story.append(Paragraph("<b>2. Immediate Containment Actions Taken:</b>", bold_label))
story.append(Paragraph(
    "• Batch MFH260712A was immediately quarantined in the centrifuge holding area.<br/>"
    "• A physical red 'QA HOLD' tag was affixed to centrifuge C-02 and mother liquor filtrate receivers.<br/>"
    "• Instrumentation maintenance was notified; valve positioner calibration confirmed mechanical jamming.<br/>"
    "• Composite samples were sent to the QC Laboratory for accelerated impurity profiling and crystal polymorph testing.",
    body_style
))
story.append(Spacer(1, 10))

# Risk & Impact assessment
story.append(Paragraph("<b>3. Quality Assurance Initial Impact Assessment:</b>", bold_label))
story.append(Paragraph(
    "Initial impact is classified as Major. Thermal excursion above 40°C during crystallization may promote unwanted degradation "
    "byproducts (specifically Related Substance Impurity-A / cyanoguanidine) and potentially alter the particle size distribution (PSD) and powder bulk density. "
    "Batch release is blocked pending full chromatographic purity re-testing and root cause investigation.",
    body_style
))

doc.build(story)
print("PDF generated successfully:", pdf_path)


# ── 2. Create TXT Incident Email ────────────────────────────────────────────
txt_path = "sample_documents/deviation_email_incident.txt"
with open(txt_path, "w", encoding="utf-8") as f:
    f.write("""FROM: r.sharma@vasudhapharma.com
TO: qa-deviations@vasudhapharma.com
DATE: 22-09-2024
SUBJECT: URGENT: Process Parameter Excursion - Crystallization Reactor R-204

Dear QA Team,

Please log an urgent manufacturing deviation for the following event:

Site / Plant: API Manufacturing Unit
Date of Occurrence: 22-09-2024
Source: Manufacturing
Related Product: Amoxicillin Trihydrate API
Batch Number: AMX-240922-04
Affected Quantity: 150 kg wet cake

Detailed Description:
During the final precipitation and vacuum filtration stage in Reactor R-204, the agitator drive motor tripped due to an overload relay fault at 03:15 AM. Agitation remained stopped for 35 minutes while slurry was in an unmixed state, causing localized agglomeration and potential non-uniform crystal distribution. The agitator was reset and run for an additional 15 minutes to re-suspend the product before filtration.

Initial Impact: Moderate
Initial Severity: Medium

Please initiate QA risk assessment and advise on sampling protocol for assay and particle size testing.

Best regards,
Ramesh Sharma
Production Manager - API Unit 1
Vasudha Pharma Chem Ltd.
""")
print("TXT email generated successfully:", txt_path)


# ── 3. Create DOCX Batch Report ─────────────────────────────────────────────
docx_path = "sample_documents/deviation_batch_excursion.docx"
doc_word = docx.Document()
doc_word.add_heading("VASUDHA PHARMA CHEM LIMITED", level=0)
doc_word.add_heading("Out of Specification / Process Deviation Record", level=1)

p = doc_word.add_paragraph()
p.add_run("Site / Plant: ").bold = True
p.add_run("QC Lab\n")
p.add_run("Date of Occurrence: ").bold = True
p.add_run("18-09-2024\n")
p.add_run("Source: ").bold = True
p.add_run("QC Lab\n")
p.add_run("Related Product: ").bold = True
p.add_run("Paracetamol API (BP/USP)\n")
p.add_run("Batch / Lot Number: ").bold = True
p.add_run("PCM-2024-B88\n")
p.add_run("Title: ").bold = True
p.add_run("OOS Result for Assay by HPLC in Batch PCM-2024-B88\n")

doc_word.add_heading("Event Details:", level=2)
doc_word.add_paragraph(
    "During routine finished API release testing of Paracetamol API batch PCM-2024-B88, "
    "the HPLC assay result obtained was 97.4% w/w on anhydrous basis against the release specification of 99.0% - 101.0% w/w. "
    "A confirmatory re-injection from the same vial yielded 97.5% w/w. System suitability criteria were met. "
    "Initial impact is Major and initial severity is High. Phase I laboratory investigation has been initiated to rule out analyst and instrument error."
)

doc_word.save(docx_path)
print("DOCX generated successfully:", docx_path)
